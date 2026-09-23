"""Tests for the finance module (spec §15): expenses, refunds, cash-flow exports."""

import datetime
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

from apps.fees.models import FeeHead, FeePayment, PaymentMethod
from apps.finance.models import Expense, ExpenseCategory, Refund
from apps.students.models import Student

U = get_user_model()


class FinanceModelsTests(TestCase):
    def test_expense_str(self):
        cat = ExpenseCategory.objects.create(name="Utilities")
        e = Expense.objects.create(
            category=cat, title="Electricity bill", amount=Decimal("5000"),
            expense_date=datetime.date(2026, 3, 1),
        )
        self.assertIn("Electricity bill", str(e))
        self.assertEqual(e.category.name, "Utilities")

    def test_expense_category_unique(self):
        ExpenseCategory.objects.create(name="Dupe")
        with self.assertRaises(Exception):
            ExpenseCategory.objects.create(name="Dupe")


class FinanceExportTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = U.objects.create_superuser("admin-finance", "a@test.test", "pw")
        cls.cat = ExpenseCategory.objects.create(name="Stationery")
        cls.expense = Expense.objects.create(
            category=cls.cat, title="Notebooks", amount=Decimal("2500"),
            expense_date=datetime.date.today(),
        )
        cls.student = Student.objects.create(
            admission_number="FIN-1", first_name="Cash", last_name="Flow"
        )
        cls.refund = Refund.objects.create(
            student=cls.student, amount=Decimal("300"), reason="Overpayment",
            refund_date=datetime.date.today(),
        )
        cls.fee = FeeHead.objects.create(name="Tuition", amount=Decimal("1000"))
        cls.payment = FeePayment.objects.create(
            student=cls.student, fee_head=cls.fee, amount=Decimal("1000"),
            method=PaymentMethod.CASH,
        )

    def setUp(self):
        self.client.force_login(self.user)

    def test_expense_export_xlsx(self):
        resp = self.client.get(reverse("finance:expense_export"))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        from openpyxl import load_workbook
        from io import BytesIO

        wb = load_workbook(BytesIO(resp.content))
        sheet = wb.active
        titles = [row[1].value for row in sheet.iter_rows(min_row=2)]
        self.assertIn("Notebooks", titles)

    def test_expense_export_csv(self):
        resp = self.client.get(reverse("finance:expense_export"), {"format": "csv"})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp["Content-Type"], "text/csv; charset=utf-8")
        self.assertIn("Notebooks".encode(), resp.content)

    def test_refund_export_csv(self):
        resp = self.client.get(reverse("finance:refund_export"), {"format": "csv"})
        self.assertEqual(resp.status_code, 200)
        self.assertIn(b"Overpayment", resp.content)

    def test_cashflow_export_net_math(self):
        resp = self.client.get(reverse("finance:cashflow_export"), {"format": "csv"})
        self.assertEqual(resp.status_code, 200)
        lines = resp.content.decode("utf-8").splitlines()
        this_month = datetime.date.today().strftime("%Y-%m")
        row = next(
            line for line in lines[1:]
            if line.split(",")[0] == this_month
        )
        parts = row.split(",")
        income, expenses, refunds, net = Decimal(parts[1]), Decimal(parts[2]), Decimal(parts[3]), Decimal(parts[4])
        self.assertEqual(net, income - expenses - refunds)

    def test_expense_create_round_trip(self):
        resp = self.client.post(
            reverse("finance:expense_add"),
            {
                "title": "New printer ink",
                "amount": "1500",
                "expense_date": datetime.date.today().isoformat(),
                "method": PaymentMethod.CASH,
            },
        )
        self.assertEqual(resp.status_code, 302)
        self.assertTrue(Expense.objects.filter(title="New printer ink").exists())